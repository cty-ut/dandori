"""Excel 连接器：读取（.xlsx、.xlsm）（语法设计 19）

原则：不悄悄地猜。
- 有多张表时必须指定表（不默认读第一张：有人调整了表的顺序，程序就会悄悄读错）
- 表头前的标题行、空行自动跳过，并告诉调用者表头在第几行（「説明」中显示）
- 公式读取 Excel 保存的计算结果；没有结果（从未用 Excel 保存过）时报错，不当作空欄

单元格先转成与 CSV 相同的文字，再交给 CSV 的类型推断（csv.build_table）：
两种文件的规则（编码保留开头的 0、金額、日付、空欄）完全一致。
- 数：按 Excel 显示的 15 位有效数字转换（0.1+0.2 的结果显示为 0.3，不是 0.30000000000000004）
- 格式带 ¥、円 的数 → 金額
- 日期 → 日付；带时刻的日期、时刻 → 文字列（段取り没有时刻的类型，不悄悄丢掉时刻）
- 合并单元格：只有左上角有值，其余是空欄
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from decimal import Decimal
from datetime import date, datetime, time
from pathlib import Path

from dandori import dates
from dandori.connectors import ConnectorError
from dandori.connectors.csv import build_table
from dandori.frontend.normalize import normalize
from dandori.table.interface import Table

SUFFIXES = (".xlsx", ".xlsm")
_YEN = ("¥", "￥", "円", "[$¥", "[$￥")


@dataclass(frozen=True)
class ExcelTable:
    table: Table
    sheet: str
    header_row: int          # 表头在第几行（Excel 的行号）


def read_excel(path: Path, display_name: str, sheet: str | None) -> ExcelTable:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    if not path.is_file():
        raise ConnectorError("E3019", file=display_name)
    try:
        values_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
        raise ConnectorError("E3074", file=display_name) from None
    try:
        name = _choose_sheet(values_book.sheetnames, sheet, display_name)
        cells = _cells(values_book[name], formula_book[name], display_name)
    finally:
        values_book.close()
        formula_book.close()
    return _table(cells, display_name, name)


def _choose_sheet(names: list[str], wanted: str | None, display_name: str) -> str:
    listing = "、".join(f"「{n}」" for n in names)
    if wanted is None:
        if len(names) == 1:
            return names[0]
        raise ConnectorError("E3071", file=display_name, count=str(len(names)), sheets=listing)
    if wanted in names:
        return wanted
    # 「１０月」和「10月」是同一张表（程序中的全角、半角）；两张表只差全角半角时不猜
    matches = [n for n in names if normalize(n) == normalize(wanted)]
    if len(matches) == 1:
        return matches[0]
    raise ConnectorError("E3072", file=display_name, sheet=wanted, sheets=listing)


def _cells(values_sheet, formula_sheet, display_name: str) -> list[tuple[int, list[str]]]:
    """每一行：(Excel 的行号, 各单元格的文字)。完全空白的行不要"""
    # 只读模式按文件中记录的范围（dimension）读取；有的工具写出的范围不对 → 按实际内容重新计算
    values_sheet.reset_dimensions()
    formula_sheet.reset_dimensions()
    rows = []
    for row_no, (values, formulas) in enumerate(
            zip(values_sheet.iter_rows(), formula_sheet.iter_rows()), start=1):
        texts = []
        for cell, formula in zip(values, formulas):
            if cell.value is None and isinstance(formula.value, str) and formula.value.startswith("="):
                raise ConnectorError("E3073", file=display_name, cell=formula.coordinate)
            texts.append(_text(cell.value, getattr(cell, "number_format", "General")))
        while texts and texts[-1] == "":
            texts.pop()
        if texts:
            rows.append((row_no, texts))
    return rows


def _text(value, number_format: str) -> str:
    """单元格 → 与 CSV 相同的文字"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return dates.format_csv(value.date())
        return f"{dates.format_csv(value.date())} {value:%H:%M:%S}"   # 带时刻：文字列
    if isinstance(value, date):
        return dates.format_csv(value)
    if isinstance(value, time):
        return f"{value:%H:%M:%S}"
    if isinstance(value, (int, float)):
        text = _number(value)
        if any(mark in (number_format or "") for mark in _YEN):
            text += "円"
        return text
    return str(value).strip()


def _number(value: int | float) -> str:
    if isinstance(value, int):
        return str(value)
    text = format(value, ".15g")                            # Excel 显示的精度
    if "e" in text:                                         # 1e-07、1e+20
        text = format(Decimal(text), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def _table(rows: list[tuple[int, list[str]]], display_name: str, sheet: str) -> ExcelTable:
    if not rows:
        raise ConnectorError("E3028", file=display_name)
    header_index = _header_index(rows)
    header_row, header = rows[header_index]
    data = []
    for row_no, cells in rows[header_index + 1:]:
        if len(cells) > len(header):
            raise ConnectorError("E3021", file=display_name, row=str(row_no))
        data.append((row_no, cells + [""] * (len(header) - len(cells))))
    table = build_table(header, data, display_name, f"{display_name}（{sheet}）")
    return ExcelTable(table, sheet, header_row)


def _header_index(rows: list[tuple[int, list[str]]]) -> int:
    """表头：跳过开头只有一格有内容的行（标题），但只在后面有更宽的行时才跳过。

        1 │ 請求一覧（2026年10月）     ← 标题：跳过
        3 │ 請求番号 │ 取引先 │ 金額  ← 表头

    只有一列的表（如会社休日）没有更宽的行：第一行就是表头。
    标题有两格以上时（「請求一覧」「2026年10月」分在两格）会被当作表头，
    之后的行比它宽，读取时报错（E3021），不会悄悄读错。
    """
    def filled(cells):
        return sum(1 for c in cells if c != "")

    if max(filled(cells) for _, cells in rows) <= 1:
        return 0
    return next(index for index, (_, cells) in enumerate(rows) if filled(cells) > 1)


# ---------- 写出（语法设计 19.3） ----------

# Excel 的表名：31 个字符以内，不能用这些字符
_SHEET_FORBIDDEN = set("[]:*?/\\")
YEN_FORMAT = '"¥"#,##0'          # 金額：读回来时仍是金額（格式中有 ¥）
DATE_FORMAT = "yyyy/mm/dd"


def check_sheet_name(name: str) -> None:
    if not name or len(name) > 31 or any(c in _SHEET_FORBIDDEN for c in name) or name.startswith("'"):
        raise ConnectorError("E3077", sheet=name)


def default_sheet(file_name: str) -> str:
    """不写表名时：文件名去掉扩展名（「出力/集計.xlsx」→「集計」）"""
    from pathlib import PurePosixPath
    return PurePosixPath(file_name.replace("\\", "/")).stem


def write_excel(sheets: list[tuple[str, Table]]) -> bytes:
    """表 → .xlsx 的字节。一个文件可以有多张表（按写出的顺序）

    写出真正的数和日期（打开后可以直接用 Excel 计算），不是文字：
    - 金額：「¥#,##0」格式；数：常规格式；日付：「yyyy/mm/dd」
    - 「00123」这样的编码是文字列，保持文字
    - 表头加粗、冻结；列宽按内容调整
    """
    import io

    import openpyxl
    from openpyxl.styles import Font

    from dandori.money import Money
    from dandori.table.interface import BLANK
    from dandori.text import display_width

    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, table in sheets:
        sheet = book.create_sheet(name)
        sheet.append(list(table.display_names))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        widths = [display_width(str(n)) for n in table.display_names]
        for row in table.rows:
            values = []
            for i, value in enumerate(row):
                if value is BLANK:
                    values.append(None)
                    continue
                if isinstance(value, Money):
                    values.append(value.amount)
                    shown = f"¥{value.amount:,}"
                else:
                    values.append(value)
                    shown = dates.format_csv(value) if isinstance(value, date) else str(value)
                widths[i] = max(widths[i], display_width(shown))
            sheet.append(values)
        for i, (_, column_type) in enumerate(table.type.columns, start=1):
            letter = openpyxl.utils.get_column_letter(i)
            fmt = {"金額": YEN_FORMAT, "日付": DATE_FORMAT}.get(column_type.value)
            if fmt:
                for (cell,) in sheet.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.number_format = fmt
            sheet.column_dimensions[letter].width = min(widths[i - 1] + 2, 60)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
