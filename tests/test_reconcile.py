"""阶段 5：入金消込"""

from datetime import date
from pathlib import Path

import pytest

from dandori.api import compile_source, run
from dandori.frontend.source import Source
from dandori.runtime.interpreter import Context

EXAMPLE = Path(__file__).parent.parent / "examples" / "入金消込"

SETTINGS = """
「入金.csv」を読み込んで、入金とする。
「請求.csv」を読み込んで、請求とする。
消込の名義は、入金の振込依頼人と請求のフリガナとする。
消込の金額は、入金の金額と請求の請求金額とする。
"""


def write(directory: Path, deposits: str, invoices: str):
    (directory / "入金.csv").write_text("振込依頼人,金額\n" + deposits, encoding="utf-8")
    (directory / "請求.csv").write_text("番号,フリガナ,請求金額,支払期日\n" + invoices, encoding="utf-8")


def reconcile(directory: Path, extra: str = "", today=date(2026, 11, 2)):
    output = []
    code = SETTINGS + extra + """
        入金を請求に消し込んで、結果とする。
        結果のそれぞれについて、
          判定を表示する。
          番号を表示する。
          理由を表示する。
        ここまで。
    """
    errors = run(Source(code), Context(output=output.append, base_dir=directory, today=today))
    assert errors == [], [e.message for e in errors]
    return [output[i:i + 3] for i in range(0, len(output), 3)]     # 每行：[判定, 番号, 理由]


# ---------- 判定的步骤 ----------

def test_exact_match(tmp_path):
    write(tmp_path, "ｶ)ﾔﾏﾀﾞｼﾖｳｼﾞ,1000\n", "A1,ヤマダショウジ,1000,2026/10/31\n")
    assert reconcile(tmp_path) == [["自動消込", "A1", "名義・金額が一致"]]


def test_fee_within_limit_is_auto(tmp_path):
    write(tmp_path, "ﾔﾏﾀﾞ,9560\n", "A1,ヤマダ,10000,2026/10/31\n")
    result = reconcile(tmp_path, "消込の手数料は、880円までとする。")
    assert result == [["自動消込", "A1", "名義が一致、手数料差引（440円）"]]


def test_fee_without_setting_needs_confirmation(tmp_path):
    # 没有写「消込の手数料」：手续费差额一律要確認（不悄悄决定）
    write(tmp_path, "ﾔﾏﾀﾞ,9560\n", "A1,ヤマダ,10000,2026/10/31\n")
    result = reconcile(tmp_path)
    assert result[0][0] == "要確認"
    assert result[1][:2] == ["未入金", "A1"]


def test_fee_over_limit_needs_confirmation(tmp_path):
    write(tmp_path, "ﾔﾏﾀﾞ,9000\n", "A1,ヤマダ,10000,2026/10/31\n")
    assert reconcile(tmp_path, "消込の手数料は、880円までとする。")[0][0] == "要確認"


def test_overpayment_is_not_fee(tmp_path):
    # 入金比请求多：不是手续费
    write(tmp_path, "ﾔﾏﾀﾞ,10440\n", "A1,ヤマダ,10000,2026/10/31\n")
    assert reconcile(tmp_path, "消込の手数料は、880円までとする。")[0][0] == "要確認"


def test_combined_payment_needs_confirmation(tmp_path):
    write(tmp_path, "ｻﾄｳ,300\n", "A1,サトウ,100,2026/10/31\nA2,サトウ,200,2026/11/30\nA3,サトウ,999,2026/12/31\n")
    result = reconcile(tmp_path, "消込の期日は、請求の支払期日とする。")
    assert [r[:2] for r in result] == [["要確認", "A1"], ["要確認", "A2"], ["未入金", "A3"]]
    assert "まとめ入金の可能性" in result[0][2]


def test_combined_payment_with_fee(tmp_path):
    write(tmp_path, "ｻﾄｳ,2560\n", "A1,サトウ,1000,2026/10/31\nA2,サトウ,2000,2026/11/30\n")
    result = reconcile(tmp_path, "消込の期日は、請求の支払期日とする。消込の手数料は、880円までとする。")
    assert [r[:2] for r in result] == [["要確認", "A1"], ["要確認", "A2"]]


def test_name_only(tmp_path):
    write(tmp_path, "ﾀｶﾊｼ,200\n", "A1,タカハシ,220,2026/10/31\n")
    result = reconcile(tmp_path)
    assert result[0] == ["要確認", "（空欄）", "名義は一致するが、金額が一致しません（候補：220円）"]


def test_amount_only(tmp_path):
    write(tmp_path, "ﾐﾔｻﾞｷ,88\n", "A1,スズキ,88,2026/10/31\n")
    result = reconcile(tmp_path)
    assert result[0] == ["要確認", "（空欄）", "金額のみ一致（候補：スズキ 88円）"]


def test_unknown_and_unpaid(tmp_path):
    write(tmp_path, "ﾀﾅｶ,50\n", "A1,スズキ,88,2026/10/31\nA2,スズキ,77,2026/12/31\n")
    result = reconcile(tmp_path, "消込の期日は、請求の支払期日とする。")
    assert result[0][:1] == ["入金不明"]
    assert result[1] == ["未入金", "A1", "対応する入金がありません（支払期日を過ぎています）"]
    assert result[2] == ["未入金", "A2", "対応する入金がありません"]          # 还没到期


def test_fifo_for_same_amounts(tmp_path):
    # 期日晚的写在前面：按期日从早到晚，先消 A2
    write(tmp_path, "ｻﾄｳ,66\n", "A1,サトウ,66,2026/11/30\nA2,サトウ,66,2026/10/31\n")
    result = reconcile(tmp_path, "消込の期日は、請求の支払期日とする。")
    assert result[0][:2] == ["自動消込", "A2"]
    assert "支払期日の早いもの" in result[0][2]
    assert result[1][:2] == ["未入金", "A1"]


def test_without_due_setting_uses_row_order(tmp_path):
    write(tmp_path, "ｻﾄｳ,66\n", "A1,サトウ,66,2026/11/30\nA2,サトウ,66,2026/10/31\n")
    assert reconcile(tmp_path)[0][:2] == ["自動消込", "A1"]


def test_invoice_used_only_once(tmp_path):
    write(tmp_path, "ｻﾄｳ,66\nｻﾄｳ,66\n", "A1,サトウ,66,2026/10/31\n")
    result = reconcile(tmp_path)
    # 第二笔：唯一的请求书已经被第一笔消掉了，没有可对应的请求书
    assert result == [["自動消込", "A1", "名義・金額が一致"],
                      ["入金不明", "（空欄）", "名義・金額とも一致する請求がありません"]]


def test_blank_amount(tmp_path):
    # 只有空欄的列会被推断为文字列，所以放一行有值的入金
    write(tmp_path, "ﾔﾏﾀﾞ,\nｽｽﾞｷ,5\n", "A1,ヤマダ,1,2026/10/31\n")
    assert reconcile(tmp_path)[0] == ["入金不明", "（空欄）", "金額が空欄です"]


def test_alias_file(tmp_path):
    write(tmp_path, "ﾀﾅｶ ｲﾁﾛｳ,33\n", "A1,タナカショウテン,33,2026/10/31\n")
    (tmp_path / "名義対応.csv").write_text("振込名義,請求側\nﾀﾅｶ ｲﾁﾛｳ,タナカショウテン\n", encoding="utf-8")
    result = reconcile(tmp_path, "消込の名義対応は「名義対応.csv」とする。")
    assert result[0][:2] == ["自動消込", "A1"]


# ---------- 结果表 ----------

def test_result_columns(tmp_path):
    write(tmp_path, "ﾔﾏﾀﾞ,1000\n", "A1,ヤマダ,1000円,2026/10/31\n")
    result = compile_source(Source(SETTINGS + "入金を請求に消し込んで、結果とする。"), base_dir=tmp_path)
    assert result.errors == []
    columns = result.program.sentences[-1].steps[0].result_type.names
    assert columns == ["判定", "理由", "振込依頼人", "金額", "番号", "フリガナ", "請求金額", "支払期日", "差額"]


def test_column_name_collision(tmp_path):
    (tmp_path / "入金.csv").write_text("振込依頼人,金額\nﾔﾏﾀﾞ,1000\n", encoding="utf-8")
    (tmp_path / "請求.csv").write_text("フリガナ,金額\nヤマダ,1000\n", encoding="utf-8")
    code = """
        「入金.csv」を読み込んで、入金とする。
        「請求.csv」を読み込んで、請求とする。
        消込の名義は、入金の振込依頼人と請求のフリガナとする。
        消込の金額は、入金の金額と請求の金額とする。
        入金を請求に消し込んで、結果とする。
        結果の「請求の金額」を表示する。
    """
    output = []
    errors = run(Source(code), Context(output=output.append, base_dir=tmp_path))
    assert errors == [] and output == ["1000"]


# ---------- 设定与错误 ----------

def compile_codes(directory, code):
    return [e.code for e in compile_source(Source(code), base_dir=directory).errors]


@pytest.mark.parametrize("extra, expected", [
    ("消込の色は、赤とする。", ["E2022"]),
    ("消込の手数料は、880円とする。", ["E3055"]),                    # 缺少「まで」
    ("消込の名義は、入金の振込依頼人とする。", ["E3055"]),            # 只有一个列
    ("消込の期日は、2026年10月31日とする。", ["E3055"]),
    ("消込の名義対応は、100とする。", ["E3055"]),
    ("消込の名義対応は「無い.csv」とする。", ["E3019"]),
    ("もし1が1ならば、消込の手数料は、880円までとする。ここまで。", ["E3054"]),
    ("消込の名義は、入金の振込依頼人と請求の金頭とする。", ["E3025"]),
])
def test_setting_errors(tmp_path, extra, expected):
    write(tmp_path, "ﾔﾏﾀﾞ,1\n", "A1,ヤマダ,1,2026/10/31\n")
    (tmp_path / "x").write_text("", encoding="utf-8")
    assert compile_codes(tmp_path, SETTINGS + extra) == expected


def test_alias_file_must_have_two_columns(tmp_path):
    write(tmp_path, "ﾔﾏﾀﾞ,1\n", "A1,ヤマダ,1,2026/10/31\n")
    (tmp_path / "名義対応.csv").write_text("a,b,c\n1,2,3\n", encoding="utf-8")
    assert compile_codes(tmp_path, SETTINGS + "消込の名義対応は「名義対応.csv」とする。") == ["E3056"]


def test_missing_required_setting(tmp_path):
    write(tmp_path, "ﾔﾏﾀﾞ,1\n", "A1,ヤマダ,1,2026/10/31\n")
    code = "「入金.csv」を読み込んで、入金とする。「請求.csv」を読み込んで、請求とする。入金を請求に消し込む。"
    assert compile_codes(tmp_path, code) == ["E3053"]


def test_setting_columns_must_exist_in_verb_tables(tmp_path):
    write(tmp_path, "ﾔﾏﾀﾞ,1\n", "A1,ヤマダ,1,2026/10/31\n")
    (tmp_path / "他.csv").write_text("x,y\n1,2\n", encoding="utf-8")
    code = SETTINGS + "「他.csv」を読み込んで、他とする。他を請求に消し込む。"
    assert compile_codes(tmp_path, code) == ["E3025"]


# ---------- 示例程序 ----------

def test_example_program(tmp_path):
    import shutil
    for name in ("入金明細.csv", "請求一覧.xlsx", "名義対応.csv", "入金消込.dandori"):
        shutil.copy(EXAMPLE / name, tmp_path / name)
    output = []
    code = (tmp_path / "入金消込.dandori").read_text(encoding="utf-8")
    errors = run(Source(code), Context(output=output.append, base_dir=tmp_path, today=date(2026, 11, 2)))
    assert errors == []
    judgements = [line.split()[0] for line in output[0].split("\n")[1:]]
    assert judgements == ["自動消込", "要確認", "要確認", "自動消込", "入金不明", "要確認", "自動消込", "要確認",
                          "未入金", "未入金"]
    # 銀行の入金明細は CSV、会社の請求一覧は Excel、結果は 1 つの Excel にシートごと
    import openpyxl
    book = openpyxl.load_workbook(tmp_path / "出力" / "消込結果.xlsx")
    assert book.sheetnames == ["消込済", "要確認", "督促候補", "入金不明"]
    assert [book[name].max_row - 1 for name in book.sheetnames] == [3, 4, 2, 1]


def test_failed_reconcile_does_not_cascade(tmp_path):
    # 设定出错 → 结果表的类型未知 → 逐行处理中的列名不再连锁报错
    write(tmp_path, "ﾔﾏﾀﾞ,\n", "A1,ヤマダ,1,2026/10/31\n")     # 金額列只有空欄：文字列
    codes = compile_codes(tmp_path, SETTINGS + "入金を請求に消し込んで、結果とする。"
                                    "結果のそれぞれについて、判定を表示する。ここまで。")
    assert codes == ["E3055"]
